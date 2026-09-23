from flask import Flask, render_template, request, jsonify
from openpyxl import load_workbook, Workbook
import os
import sys
import webbrowser
import threading


def ruta_base():
    """Carpeta donde vive el programa. Si es un .exe empaquetado con
    PyInstaller, es la carpeta donde está el .exe; si es el script de
    Python normal, es la carpeta donde está app.py."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def ruta_recurso(relativa):
    """Ubica archivos empaquetados dentro del .exe (como templates/),
    que PyInstaller extrae a una carpeta temporal en tiempo de ejecución."""
    base = getattr(sys, "_MEIPASS", ruta_base())
    return os.path.join(base, relativa)


app = Flask(__name__, template_folder=ruta_recurso("templates"))

# El Excel se guarda SIEMPRE junto al programa (junto al .exe o junto a
# app.py), así que funciona igual en tu máquina y en la de tu equipo,
# sin rutas fijas a un usuario en particular.
ARCHIVO_AGENCIA = os.path.join(ruta_base(), "agencia_viajes.xlsx")


# ==========================================
# HELPERS PARA TRABAJAR CON EL EXCEL
# ==========================================

def abrir_libro():
    """Abre el libro si existe; si no, crea uno nuevo y vacío (sin la
    hoja 'Sheet' por defecto, para que solo queden las hojas que nosotros
    creemos a propósito)."""
    if os.path.exists(ARCHIVO_AGENCIA):
        return load_workbook(ARCHIVO_AGENCIA)

    libro = Workbook()
    libro.remove(libro.active)
    return libro


def obtener_hoja(libro, nombre, encabezados):
    """Devuelve la hoja si ya existe; si no, la crea con sus encabezados."""
    if nombre in libro.sheetnames:
        return libro[nombre]

    hoja = libro.create_sheet(nombre)
    hoja.append(encabezados)
    return hoja


def leer_hoja(nombre, columnas):
    """Lee todas las filas de una hoja (desde la fila 2) y las devuelve
    como lista de diccionarios usando los nombres de 'columnas'."""
    if not os.path.exists(ARCHIVO_AGENCIA):
        return []

    libro = load_workbook(ARCHIVO_AGENCIA)
    if nombre not in libro.sheetnames:
        return []

    hoja = libro[nombre]
    filas = []
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if fila[0] is None:
            continue
        filas.append({columnas[i]: fila[i] for i in range(len(columnas))})

    return filas


def validar_campos(datos, requeridos):
    """Devuelve la lista de campos requeridos que faltan o vienen vacíos."""
    return [c for c in requeridos if datos.get(c) in (None, "")]


def siguiente_folio(hoja):
    """Genera un folio tipo C-001, C-002... basado en cuántas filas ya hay."""
    total = hoja.max_row  # incluye encabezado
    return f"C-{total:03d}"


def seed_alternativas(libro):
    """Si la hoja de Alternativas se está creando por primera vez, le
    metemos un par de ejemplos para que el Cotizador no se vea vacío."""
    if "Alternativas" in libro.sheetnames:
        return

    hoja = libro.create_sheet("Alternativas")
    hoja.append(["Destino", "Categoria", "Aerolinea", "Precio", "Imagen"])
    hoja.append(["Tokio", "Recomendada", "United Airlines", 8100, "🗼"])
    hoja.append(["Tokio", "VIP", "ANA All Nippon Airways", 12400, "🗼"])
    hoja.append(["Tokio", "Tarifa accesible", "Aeromexico", 6200, "🗼"])


# ==========================================
# PÁGINA PRINCIPAL
# ==========================================

@app.route("/")
def inicio():
    return render_template("index.html")


# ==========================================
# PERSONAS (clientes)
# ==========================================

@app.route("/persona", methods=["POST"])
def recibir_persona():
    datos = request.get_json(silent=True)
    if not datos:
        return jsonify({"error": "JSON inválido o vacío"}), 400

    faltantes = validar_campos(datos, ["nombre", "edad"])
    if faltantes:
        return jsonify({"error": "Faltan campos obligatorios", "campos": faltantes}), 400

    libro = abrir_libro()
    hoja = obtener_hoja(libro, "Personas", ["Nombre", "Edad"])
    hoja.append([datos["nombre"], datos["edad"]])
    libro.save(ARCHIVO_AGENCIA)

    return jsonify({"mensaje": "Persona guardada en Excel", **datos})


@app.route("/personas", methods=["GET"])
def obtener_personas():
    return jsonify(leer_hoja("Personas", ["nombre", "edad"]))


# ==========================================
# SOLICITUDES DE VIAJE
# ==========================================

@app.route("/solicitud", methods=["POST"])
def recibir_solicitud():
    datos = request.get_json(silent=True)
    if not datos:
        return jsonify({"error": "JSON inválido o vacío"}), 400

    campos = ["cliente", "destino", "categoria", "pasajeros", "salida", "regreso"]
    faltantes = validar_campos(datos, campos)
    if faltantes:
        return jsonify({"error": "Faltan campos obligatorios", "campos": faltantes}), 400

    libro = abrir_libro()
    hoja = obtener_hoja(libro, "Solicitudes",
                         ["Cliente", "Destino", "Categoria", "Pasajeros", "Salida", "Regreso"])
    hoja.append([datos[c] for c in campos])
    libro.save(ARCHIVO_AGENCIA)

    return jsonify({"mensaje": "Solicitud guardada correctamente"})


@app.route("/solicitudes", methods=["GET"])
def obtener_solicitudes():
    columnas = ["cliente", "destino", "categoria", "pasajeros", "salida", "regreso"]
    return jsonify(leer_hoja("Solicitudes", columnas))


# ==========================================
# VUELOS DISPONIBLES (por ruta/destino)
# ==========================================

@app.route("/vuelo", methods=["POST"])
def recibir_vuelo():
    datos = request.get_json(silent=True)
    if not datos:
        return jsonify({"error": "JSON inválido o vacío"}), 400

    campos = ["destino", "aerolinea", "ruta", "escala", "tarifa_neta", "cupos"]
    faltantes = validar_campos(datos, campos)
    if faltantes:
        return jsonify({"error": "Faltan campos obligatorios", "campos": faltantes}), 400

    libro = abrir_libro()
    hoja = obtener_hoja(libro, "Vuelos",
                         ["Destino", "Aerolinea", "Ruta", "Escala", "Tarifa neta", "Cupos"])
    hoja.append([datos[c] for c in campos])
    libro.save(ARCHIVO_AGENCIA)

    return jsonify({"mensaje": "Vuelo guardado correctamente"})


@app.route("/vuelos", methods=["GET"])
def obtener_vuelos():
    columnas = ["destino", "aerolinea", "ruta", "escala", "tarifa_neta", "cupos"]
    vuelos = leer_hoja("Vuelos", columnas)

    destino = request.args.get("destino", "").strip().lower()
    if destino:
        vuelos = [v for v in vuelos if str(v["destino"]).strip().lower() == destino]

    return jsonify(vuelos)


# ==========================================
# HOSPEDAJES
# ==========================================

@app.route("/hospedaje", methods=["POST"])
def recibir_hospedaje():
    datos = request.get_json(silent=True)
    if not datos:
        return jsonify({"error": "JSON inválido o vacío"}), 400

    campos = ["destino", "hospedaje", "noches", "zona", "neto_total", "categoria"]
    faltantes = validar_campos(datos, campos)
    if faltantes:
        return jsonify({"error": "Faltan campos obligatorios", "campos": faltantes}), 400

    libro = abrir_libro()
    hoja = obtener_hoja(libro, "Hospedajes",
                         ["Destino", "Hospedaje", "Noches", "Zona", "Neto total", "Categoria"])
    hoja.append([datos[c] for c in campos])
    libro.save(ARCHIVO_AGENCIA)

    return jsonify({"mensaje": "Hospedaje guardado correctamente"})


@app.route("/hospedajes", methods=["GET"])
def obtener_hospedajes():
    columnas = ["destino", "hospedaje", "noches", "zona", "neto_total", "categoria"]
    hospedajes = leer_hoja("Hospedajes", columnas)

    destino = request.args.get("destino", "").strip().lower()
    if destino:
        hospedajes = [h for h in hospedajes if str(h["destino"]).strip().lower() == destino]

    return jsonify(hospedajes)


# ==========================================
# ALTERNATIVAS (catálogo: recomendada / vip / accesible)
# ==========================================

@app.route("/alternativa", methods=["POST"])
def recibir_alternativa():
    datos = request.get_json(silent=True)
    if not datos:
        return jsonify({"error": "JSON inválido o vacío"}), 400

    campos = ["destino", "categoria", "aerolinea", "precio"]
    faltantes = validar_campos(datos, campos)
    if faltantes:
        return jsonify({"error": "Faltan campos obligatorios", "campos": faltantes}), 400

    imagen = datos.get("imagen") or "🧳"

    libro = abrir_libro()
    seed_alternativas(libro)
    hoja = obtener_hoja(libro, "Alternativas", ["Destino", "Categoria", "Aerolinea", "Precio", "Imagen"])
    hoja.append([datos["destino"], datos["categoria"], datos["aerolinea"], datos["precio"], imagen])
    libro.save(ARCHIVO_AGENCIA)

    return jsonify({"mensaje": "Alternativa guardada correctamente"})


@app.route("/alternativas", methods=["GET"])
def obtener_alternativas():
    if not os.path.exists(ARCHIVO_AGENCIA):
        libro = abrir_libro()
        seed_alternativas(libro)
        libro.save(ARCHIVO_AGENCIA)
    else:
        libro = load_workbook(ARCHIVO_AGENCIA)
        if "Alternativas" not in libro.sheetnames:
            seed_alternativas(libro)
            libro.save(ARCHIVO_AGENCIA)

    columnas = ["destino", "categoria", "aerolinea", "precio", "imagen"]
    alternativas = leer_hoja("Alternativas", columnas)

    destino = request.args.get("destino", "").strip().lower()
    if destino:
        alternativas = [a for a in alternativas if str(a["destino"]).strip().lower() == destino]

    return jsonify(alternativas)


# ==========================================
# DOCUMENTACIÓN NECESARIA (checklist por destino)
# ==========================================

@app.route("/documentacion", methods=["POST"])
def guardar_documentacion():
    datos = request.get_json(silent=True)
    if not datos:
        return jsonify({"error": "JSON inválido o vacío"}), 400

    faltantes = validar_campos(datos, ["destino"])
    if faltantes:
        return jsonify({"error": "Falta el destino"}), 400

    viable = bool(datos.get("viable", True))
    fila_valores = [
        datos["destino"],
        bool(datos.get("pasaporte", False)),
        bool(datos.get("visa", False)),
        bool(datos.get("identificacion", False)),
        datos.get("otros", "") or "",
        viable,
        "" if viable else (datos.get("razon", "") or ""),
    ]

    libro = abrir_libro()
    hoja = obtener_hoja(libro, "Documentacion",
                         ["Destino", "Pasaporte", "Visa", "Identificacion", "Otros", "Viable", "Razon"])

    # Un destino = una sola fila. Si ya existe, la actualizamos en vez de duplicarla.
    destino_normalizado = str(datos["destino"]).strip().lower()
    fila_encontrada = None
    for i in range(2, hoja.max_row + 1):
        valor_celda = hoja.cell(row=i, column=1).value
        if valor_celda and str(valor_celda).strip().lower() == destino_normalizado:
            fila_encontrada = i
            break

    if fila_encontrada:
        for col, valor in enumerate(fila_valores, start=1):
            hoja.cell(row=fila_encontrada, column=col, value=valor)
    else:
        hoja.append(fila_valores)

    libro.save(ARCHIVO_AGENCIA)
    return jsonify({"mensaje": "Documentación guardada correctamente"})


@app.route("/documentacion", methods=["GET"])
def obtener_documentacion():
    columnas = ["destino", "pasaporte", "visa", "identificacion", "otros", "viable", "razon"]
    registros = leer_hoja("Documentacion", columnas)

    destino = request.args.get("destino", "").strip().lower()
    if destino:
        registros = [r for r in registros if str(r["destino"]).strip().lower() == destino]

    return jsonify(registros[0] if (destino and registros) else (None if destino else registros))




@app.route("/cotizacion", methods=["POST"])
def recibir_cotizacion():
    datos = request.get_json(silent=True)
    if not datos:
        return jsonify({"error": "JSON inválido o vacío"}), 400

    campos = ["cliente", "destino", "costo_neto", "utilidad", "total_cliente", "estado"]
    faltantes = validar_campos(datos, campos)
    if faltantes:
        return jsonify({"error": "Faltan campos obligatorios", "campos": faltantes}), 400

    libro = abrir_libro()
    hoja = obtener_hoja(libro, "Cotizaciones",
                         ["Folio", "Cliente", "Destino", "Costo neto", "Utilidad", "Total cliente", "Estado"])

    folio = datos.get("folio") or siguiente_folio(hoja)
    hoja.append([folio, datos["cliente"], datos["destino"], datos["costo_neto"],
                 datos["utilidad"], datos["total_cliente"], datos["estado"]])
    libro.save(ARCHIVO_AGENCIA)

    return jsonify({"mensaje": "Cotización guardada correctamente", "folio": folio})


@app.route("/cotizaciones", methods=["GET"])
def obtener_cotizaciones():
    columnas = ["folio", "cliente", "destino", "costo_neto", "utilidad", "total_cliente", "estado"]
    return jsonify(leer_hoja("Cotizaciones", columnas))


if __name__ == "__main__":
    # Abre el navegador solo, sin que el usuario tenga que escribir la URL
    threading.Timer(1.0, lambda: webbrowser.open("http://127.0.0.1:5000/")).start()

    app.run(debug=False)